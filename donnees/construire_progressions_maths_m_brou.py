"""Construit donnees/progressions_maths_m_brou_2026.json depuis les fichiers Word de M. Brou.

  python3 donnees/construire_progressions_maths_m_brou.py donnees "chemin/du/dossier/Word/"

Les fichiers Word ne sont pas dans le depot ; seul le texte extrait (JSON) l'est.
"""
import hashlib, json, os, re, sys
import docx
from docx.oxml.ns import qn

M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}t"
W = qn("w:t")
ENSEMBLES = {"R": "ℝ", "N": "ℕ", "Z": "ℤ", "Q": "ℚ", "C": "ℂ", "R×R": "ℝ×ℝ", "RxR": "ℝ×ℝ"}
DOSSIER = sys.argv[2] if len(sys.argv) > 2 else "Progression_Mr Brou/"
DOSSIER = DOSSIER if DOSSIER.endswith("/") else DOSSIER + "/"
# fichier, classe inscrite, niveau, serie (codes du bot), libelle
FICHIERS = [
    ("PROGRESSION 1B 2026.docx", "1B", "PREMIERE", "B", "Première B"),
    ("PROGRESSION 1E 2026.docx", "1E", "PREMIERE", "E", "Première E"),
    ("PROGRESSION 1F 1, 2, 3 2026.docx", "1F1,2,3", "PREMIERE", "F1F2F3", "Première F1, F2, F3"),
    ("PROGRESSION 1F 7 2025   .docx", "1F7", "PREMIERE", "F7", "Première F7"),
    ("PROGRESSION 1G2  2026.docx", "1G2", "PREMIERE", "G2", "Première G2"),
    ("PROGRESSION 2NDE AB 2026.docx", "2NDE AB", "SECONDE", "B", "Seconde AB"),
    ("PROGRESSION 2NDE G1 2026.docx", "2NDE G1", "SECONDE", "G1", "Seconde G1"),
    ("PROGRESSION 2NDE G2 2026.docx", "2NDE G2", "SECONDE", "G2", "Seconde G2"),
    ("PROGRESSION 2NDE T1 2026.docx", "2T1", "SECONDE", "EF1F2F3", "Seconde T1 (E, F1, F2, F3)"),
    ("PROGRESSION 2NDE T2 2026.docx", "2T2", "SECONDE", "F4", "Seconde T2 (F4)"),
    ("PROGRESSION 2NDE T3 2026.docx", "2T3", "SECONDE", "F7", "Seconde T3 (F7)"),
    ("PROGRESSION TB  2026.docx", "TB", "TERMINALE", "B", "Terminale B"),
    ("PROGRESSION TF 1, 2, 3 2026.docx", "TF1,2,3", "TERMINALE", "F1F2F3", "Terminale F1, F2, F3"),
    ("PROGRESSION TF 7  2025    .docx", "TF7", "TERMINALE", "F7", "Terminale F7"),
    ("PROGRESSION TG1 2026.docx", "TG1", "TERMINALE", "G1", "Terminale G1"),
    ("PROGRESSION TG2  2026.docx", "TG2", "TERMINALE", "G2", "Terminale G2"),
]


def texte_para(p):
    morceaux = []
    for e in p._p.iter():
        if e.tag == W:
            morceaux.append(e.text or "")
        elif e.tag == M:
            t = (e.text or "").strip()
            morceaux.append(f" {ENSEMBLES[t]} " if t in ENSEMBLES else (e.text or ""))
    texte = " ".join("".join(morceaux).split())
    return re.sub(r"\s+([;,.)])", r"\1", texte)


def texte_cellule(c, sep="; "):
    return sep.join(t for t in (texte_para(p) for p in c.paragraphs) if t)


def lignes_tableau(d):
    lignes = []
    for t in d.tables:
        for r in t.rows:
            cells, prec = [], None
            for c in r.cells:
                if c._tc is prec:
                    continue
                prec = c._tc
                # La colonne des lecons (3e) : un titre sur plusieurs lignes, pas une liste.
                cells.append(texte_cellule(c, " " if len(cells) == 2 else "; "))
            lignes.append(cells)
    return lignes


def mettre_en_forme(chemin, libelle):
    d = docx.Document(chemin)
    entete = [texte_para(p) for p in d.paragraphs if texte_para(p)]
    titre = next((e for e in entete if e.upper().startswith("PROGRESSION")), "")
    import zipfile
    brut = " ".join(re.sub(r"<[^>]+>", "", zipfile.ZipFile(chemin).read(n).decode("utf8", "ignore"))
                    for n in zipfile.ZipFile(chemin).namelist() if n.startswith("word/") and n.endswith(".xml"))
    edition = re.search(r"SEPTEMBRE\s*20\d\d", brut.upper())
    lignes = lignes_tableau(d)
    horaire = next((l[0] for l in lignes if len(l) == 1 and "HEURE" in l[0].upper()), "")
    # Lignes : mois, semaine, lecon, contenus, volume horaire. Les semaines consecutives
    # d'une meme lecon sont regroupees.
    groupes = []
    mois_courant = ""
    for l in lignes:
        if len(l) < 4 or l[0].upper().startswith("MOIS"):
            continue
        mois, sem, lecon, contenus = l[0], l[1], l[2], l[3]
        vh = l[4] if len(l) > 4 else ""
        mois = mois.capitalize() if mois else mois_courant
        mois_courant = mois
        if not lecon and not contenus:
            continue
        cle = (lecon, contenus)
        if groupes and groupes[-1]["cle"] == cle:
            g = groupes[-1]
            if mois and mois not in g["mois"]:
                g["mois"].append(mois)
            if sem and sem not in g["sem"]:
                g["sem"].append(sem)
            continue
        groupes.append({"cle": cle, "mois": [mois] if mois else [], "sem": [sem] if sem else [], "vh": vh})
    sortie = [f"{titre} - Mathématiques - {libelle}".strip(" -"),
              "Ministère de l'Éducation nationale, de l'Alphabétisation et de l'Enseignement technique - Enseignement technique (MET-FPA)"]
    if edition:
        sortie.append(f"Édition : {edition.group(0).capitalize()}")
    if horaire:
        sortie.append(f"Volume horaire : {horaire}")
    sortie += [e for e in entete if e.upper().startswith(("N.B", "LE CONTENU"))]
    sortie.append("")
    for g in groupes:
        lecon, contenus = g["cle"]
        sem = g["sem"]
        quand = "/".join(g["mois"])
        if sem:
            quand += f", semaine{'s' if len(sem) > 1 else ''} {sem[0]}{'-' + sem[-1] if len(sem) > 1 else ''}"
        ligne = f"{quand} : {lecon or '(suite)'}"
        if g["vh"]:
            ligne += f" ({g['vh']})"
        if contenus:
            ligne += f" - {contenus}"
        sortie.append(ligne)
    return "\n".join(sortie), (edition.group(0).title() if edition else "")


docs = []
for fichier, classe, niveau, serie, libelle in FICHIERS:
    chemin = DOSSIER + fichier
    contenu = open(chemin, "rb").read()
    texte, edition = mettre_en_forme(chemin, libelle)
    docs.append({
        "fichier": fichier.strip(), "classe_source": classe, "niveau": niveau, "serie": serie,
        "libelle": libelle, "edition_source": edition, "sha256": hashlib.sha256(contenu).hexdigest(),
        "texte": texte,
    })
os.makedirs(sys.argv[1], exist_ok=True)
with open(os.path.join(sys.argv[1], "progressions_maths_m_brou_2026.json"), "w", encoding="utf-8") as f:
    json.dump({"transmis_par": "M. Brou, enseignant de mathématiques", "recu_le": "2026-10-04",
               "documents": docs}, f, ensure_ascii=False, indent=1)
for d in docs:
    print(d["classe_source"], d["niveau"], d["serie"], d["edition_source"], len(d["texte"]))
